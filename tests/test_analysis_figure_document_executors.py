from __future__ import annotations

import csv
import importlib.util
import json
import math
import sys
import tempfile
import unittest
from pathlib import Path

import numpy as np
import jsonschema
from scipy import optimize, special, stats


ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "evals" / "fixtures" / "real-artifacts"
for relative in (
    "skills/zju-statistics-audit/scripts",
    "skills/zju-scientific-figure/scripts",
    "skills/zju-scientific-writing/scripts",
    "skills/zju-paper2ppt/scripts",
    "skills/zju-proposal-writer/scripts",
):
    value = str(ROOT / relative)
    if value not in sys.path:
        sys.path.insert(0, value)


def load_module(name: str, relative: str):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


class AnalysisFigureDocumentExecutorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.analysis = load_module(
            "l2_execute_analysis",
            "skills/zju-statistics-audit/scripts/execute_analysis.py",
        )
        cls.figure = load_module(
            "l2_render_figure",
            "skills/zju-scientific-figure/scripts/render_from_registry.py",
        )
        cls.docx = load_module(
            "l2_build_docx",
            "skills/zju-scientific-writing/scripts/build_research_docx.py",
        )
        cls.pptx = load_module(
            "l2_build_pptx",
            "skills/zju-paper2ppt/scripts/build_presentation.py",
        )
        cls.proposal = load_module(
            "l2_render_proposal",
            "skills/zju-proposal-writer/scripts/render_proposal_docx.py",
        )

    def test_real_xlsx_profile_and_one_way_anova_match_scipy(self):
        rows = self.analysis.load_rows(FIXTURES / "analysis-input.xlsx")
        profile = self.analysis.profile_rows(rows)
        self.assertEqual(profile["rows"], 18)
        self.assertEqual(profile["columns"], 3)
        self.assertEqual(profile["duplicate_full_rows"], 0)
        registry, run = self.analysis.execute(
            FIXTURES / "analysis-input.xlsx",
            FIXTURES / "analysis-contract.json",
        )
        result = registry["results"][0]
        grouped = [
            np.asarray([float(row["response"]) for row in rows if row["group"] == group])
            for group in ("control", "low", "high")
        ]
        expected = stats.f_oneway(*grouped)
        grand = np.concatenate(grouped)
        eta = sum(len(group) * (group.mean() - grand.mean()) ** 2 for group in grouped) / np.sum((grand - grand.mean()) ** 2)
        self.assertAlmostEqual(result["test_statistic"], float(expected.statistic), places=12)
        self.assertAlmostEqual(result["p_value"], float(expected.pvalue), places=12)
        self.assertAlmostEqual(result["estimate"], float(eta), places=12)
        self.assertLessEqual(result["ci"]["lower"], result["estimate"])
        self.assertGreaterEqual(result["ci"]["upper"], result["estimate"])
        self.assertEqual(result["n"], {"experimental_units": 18, "observations": 18})
        self.assertEqual(run["dataset_profile"]["rows"], 18)

    def test_binomial_glm_matches_independent_scipy_optimization(self):
        x = np.linspace(-2.9, 2.9, 40)
        y = np.asarray([
            int(((index * 37) % 100) / 100 < special.expit(-0.25 + 0.9 * value))
            for index, value in enumerate(x)
        ])
        rows = [
            {"unit": f"U{index:02d}", "dose": float(value), "event": "yes" if outcome else "no"}
            for index, (value, outcome) in enumerate(zip(x, y), 1)
        ]
        contract = {
            "contract_id": "AC-GLM-1",
            "study_id": "STUDY-GLM-1",
            "design_stage": "frozen",
            "design": {"grouping_structure": "one independent row per unit"},
            "outcomes": [{"outcome_id": "OUT-1"}],
            "analyses": [{
                "analysis_id": "AN-GLM",
                "outcome_ids": ["OUT-1"],
                "analysis_population": "all complete independent units",
                "effect_measure": "log_odds_ratio_per_unit",
                "multiplicity_family": "primary_single_test",
                "execution": {
                    "method": "binomial_logistic_glm",
                    "result_id": "RES-GLM",
                    "outcome_column": "event",
                    "predictor_column": "dose",
                    "experimental_unit_column": "unit",
                    "event_value": "yes",
                    "non_event_value": "no",
                },
            }],
        }
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            data_path = base / "glm.csv"
            contract_path = base / "contract.json"
            write_csv(data_path, rows)
            contract_path.write_text(json.dumps(contract), encoding="utf-8")
            registry, _ = self.analysis.execute(data_path, contract_path)
        result = registry["results"][0]
        matrix = np.column_stack([np.ones(len(x)), x])

        def objective(beta):
            eta = matrix @ beta
            return float(np.sum(np.logaddexp(0.0, eta) - y * eta))

        reference = optimize.minimize(objective, np.zeros(2), method="BFGS", options={"gtol": 1e-11})
        self.assertAlmostEqual(result["intercept"], float(reference.x[0]), places=7)
        self.assertAlmostEqual(result["estimate"], float(reference.x[1]), places=7)
        self.assertAlmostEqual(result["odds_ratio"], math.exp(float(reference.x[1])), places=7)
        self.assertEqual(result["verification_scope"], "independent_units_one_continuous_predictor_unweighted_binomial_glm_only")

    def test_unsupported_dependence_and_separation_fail_closed(self):
        rows = [
            {"unit": f"U{index:02d}", "dose": float(index), "event": "no" if index < 20 else "yes"}
            for index in range(40)
        ]
        contract = {
            "contract_id": "AC-SEP",
            "study_id": "STUDY-SEP",
            "design_stage": "frozen",
            "design": {"grouping_structure": "independent"},
            "outcomes": [{"outcome_id": "OUT"}],
            "analyses": [{
                "analysis_id": "AN",
                "outcome_ids": ["OUT"],
                "analysis_population": "complete units",
                "effect_measure": "log_odds_ratio_per_unit",
                "multiplicity_family": "single",
                "execution": {
                    "method": "binomial_logistic_glm",
                    "outcome_column": "event",
                    "predictor_column": "dose",
                    "experimental_unit_column": "unit",
                    "event_value": "yes",
                    "non_event_value": "no",
                },
            }],
        }
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            data_path = base / "data.csv"
            contract_path = base / "contract.json"
            write_csv(data_path, rows)
            contract_path.write_text(json.dumps(contract), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "separation|converge"):
                self.analysis.execute(data_path, contract_path)
            contract["design"]["clusters"] = "laboratory"
            contract_path.write_text(json.dumps(contract), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "clustered or nested"):
                self.analysis.execute(data_path, contract_path)

    def test_registry_figure_exports_pass_data_and_structure_qa(self):
        registry, run = self.analysis.execute(
            FIXTURES / "analysis-input.xlsx",
            FIXTURES / "analysis-contract.json",
        )
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            paths = self.analysis.write_outputs(registry, run, base)
            manifest, manifest_path = self.figure.render(
                FIXTURES / "analysis-input.xlsx",
                paths["registry"],
                FIXTURES / "figure-spec.json",
                base / "figure",
            )
            self.assertTrue(manifest_path.is_file())
            self.assertEqual(manifest["artifact_qa"]["data_to_mark_audit"]["status"], "passed")
            self.assertEqual(manifest["artifact_qa"]["data_to_mark_audit"]["rendered_observations"], 18)
            self.assertEqual({item["format"] for item in manifest["exports"]}, {"png", "svg", "pdf"})
            self.assertTrue(all(item["machine_check"]["status"] == "passed" for item in manifest["exports"]))

    def test_docx_and_pptx_builders_reopen_and_resolve_sources(self):
        registry_path = FIXTURES / "analysis-output" / "result-registry.json"
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            docx_path = base / "report.docx"
            doc_manifest = self.docx.build_docx(
                FIXTURES / "research-document.json",
                docx_path,
                registry_path=registry_path,
            )
            self.assertEqual(doc_manifest["qa"]["ooxml_reopen"], "passed")
            self.assertTrue(self.docx.verify_docx(docx_path)["valid"])

            pptx_path = base / "deck.pptx"
            deck_manifest = self.pptx.build_presentation(
                FIXTURES / "deck-plan.json",
                pptx_path,
                registry_path=registry_path,
            )
            self.assertEqual(deck_manifest["slide_count"], 2)
            self.assertTrue(self.pptx.verify_pptx(pptx_path, expected_slide_count=2)["valid"])

            try:
                from docx import Document
            except ImportError:
                Document = None
            if Document is not None:
                self.assertEqual(Document(docx_path).paragraphs[0].text, "Dose-response analysis report")
            try:
                from pptx import Presentation
            except ImportError:
                Presentation = None
            if Presentation is not None:
                presentation = Presentation(pptx_path)
                self.assertEqual(len(presentation.slides), 2)
                self.assertIn("[Sources]", presentation.slides[0].notes_slide.notes_text_frame.text)

    def test_proposal_manifest_renders_real_docx(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "proposal.docx"
            manifest, source_path = self.proposal.render(
                FIXTURES / "proposal-manifest.json",
                output,
            )
            self.assertTrue(source_path.is_file())
            self.assertEqual(manifest["qa"]["ooxml_reopen"], "passed")
            self.assertEqual(manifest["proposal_validation"]["valid"], True)

    def test_real_binary_fixtures_are_reopenable(self):
        self.assertTrue(self.docx.verify_docx(FIXTURES / "research-document.docx")["valid"])
        self.assertTrue(self.docx.verify_docx(FIXTURES / "research-proposal.docx")["valid"])
        self.assertTrue(self.pptx.verify_pptx(
            FIXTURES / "research-deck.pptx",
            expected_slide_count=2,
        )["valid"])

    def test_executor_input_fixtures_match_published_schemas(self):
        schema_cases = [
            (
                "skills/zju-statistics-audit/references/bounded-execution.schema.json",
                json.loads((FIXTURES / "analysis-contract.json").read_text(encoding="utf-8"))["analyses"][0]["execution"],
            ),
            (
                "skills/zju-scientific-figure/references/executable-figure.schema.json",
                json.loads((FIXTURES / "figure-spec.json").read_text(encoding="utf-8")),
            ),
            (
                "skills/zju-scientific-writing/references/research-document-source.schema.json",
                json.loads((FIXTURES / "research-document.json").read_text(encoding="utf-8")),
            ),
            (
                "skills/zju-paper2ppt/references/deck-plan.schema.json",
                json.loads((FIXTURES / "deck-plan.json").read_text(encoding="utf-8")),
            ),
        ]
        for schema_path, instance in schema_cases:
            with self.subTest(schema=schema_path):
                schema = json.loads((ROOT / schema_path).read_text(encoding="utf-8"))
                jsonschema.Draft202012Validator(schema).validate(instance)


if __name__ == "__main__":
    unittest.main()
